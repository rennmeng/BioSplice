#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Per-event annotation and the main analysis pipeline.

Entry point: filter_and_summarize(rmats_dir, output_dir, genome_file, gtf_file, ...)
"""

import glob
import logging
import os
import re
import uuid
from collections import defaultdict
from datetime import datetime

import pandas as pd
from pyfaidx import Fasta
from Bio.Seq import Seq
from openpyxl.styles import Alignment

from . import console, settings
from .alt_isoform import compute_alt_from_ref
from .events import (
    ANNOTATION_COLUMNS,
    RMATS_COORD_COLS,
    RMATS_META_COLS,
    _is_frameshift_by_event,
)
from .genbank_output import generate_exon_gb, generate_tran_gb
from .gtf_index import build_gtf_index, build_gtf_index_with_cache
from .matching import anchor_ref_transcript
from .nmd import _get_start_codon_position, judge_nmd_status
from .sequences import (
    _compute_trimmed_diff,
    _starts_with_atg,
    extend_cds_with_3utr,
    get_coding_sequence,
    truncate_cds_at_stop,
)

logger = logging.getLogger(__name__)

__all__ = ["filter_and_summarize", "process_event", "ANNOTATION_COLUMNS"]


def parse_jc_single(s) -> float:
    if pd.isna(s) or str(s).strip() in ("", ".", "nan"):
        return 0.0
    try:
        vals = [float(x) for x in str(s).split(',') if x.strip()]
        return sum(vals) / len(vals) if vals else 0.0
    except Exception:
        return 0.0


def _fill_missing_gene_symbols(df: pd.DataFrame) -> pd.DataFrame:
    if 'geneSymbol' not in df.columns or 'GeneID' not in df.columns:
        return df

    mask_empty_symbol = df['geneSymbol'].isna() | (df['geneSymbol'].astype(str).str.strip() == '')
    mask_valid_geneid = df['GeneID'].notna() & (df['GeneID'].astype(str).str.strip() != '')
    fill_mask = mask_empty_symbol & mask_valid_geneid

    n_filled = fill_mask.sum()
    if n_filled > 0:
        df.loc[fill_mask, 'geneSymbol'] = df.loc[fill_mask, 'GeneID'].astype(str).str.strip()
        logger.info(f"  📝 Filled {n_filled} missing geneSymbol(s) with GeneID values")

    return df


def _build_all_significant_row(df_filt, annot_df, etype, global_counter_start=0):
    dup_cols = [c for c in df_filt.columns if re.match(r'^.*\.\d+$', c)]
    if dup_cols:
        df_filt = df_filt.drop(columns=dup_cols)
    result = pd.DataFrame(index=df_filt.index)

    n_rows = len(df_filt)
    ids = list(range(global_counter_start + 1, global_counter_start + n_rows + 1))
    result['Event_ID'] = ids
    result['Splice_Type'] = etype

    for col in RMATS_META_COLS:
        result[col] = df_filt[col] if col in df_filt.columns else ''
    for col in ANNOTATION_COLUMNS:
        result[col] = annot_df[col]
    for col in RMATS_COORD_COLS:
        result[col] = df_filt[col] if col in df_filt.columns else ''
    return result, ids


def process_event(row, etype, genome, gene_idx, gene_exons, tid_strand, alias_map,
                  longest_tx_per_gene, gb_output_dir=None, event_id="", job_id=None):
    res = {col: '' for col in ANNOTATION_COLUMNS}
    try:
        chrom = row['chr']
        gene_id = str(row.get('GeneID', row.get('geneSymbol', ''))).strip()
        if not gene_id:
            res['Match_Status'] = "[LEVEL0] ERROR: No GeneID/Symbol"
            return res

        ref_tid, ref_cds_raw, exon_bounds, strand, diag_msg = anchor_ref_transcript(
            row, etype, gene_id, gene_idx, gene_exons, tid_strand, alias_map, longest_tx_per_gene)

        if not ref_tid or not exon_bounds:
            res['Match_Status'] = diag_msg or "[LEVEL0] UNKNOWN_MATCH_FAILURE"
            return res

        res['Ref_Transcript_ID'] = ref_tid
        res['Match_Status'] = diag_msg

        raw_gene_symbol = str(row.get('geneSymbol', '')).strip()
        display_gene_symbol = raw_gene_symbol if raw_gene_symbol else re.sub(r'_\d+$', '', str(gene_id))
        canonical_gene = alias_map.get(display_gene_symbol, display_gene_symbol)

        full_tx_exons = gene_exons.get(canonical_gene, {}).get(ref_tid, exon_bounds)
        if not full_tx_exons:
            full_tx_exons = exon_bounds

        if ref_cds_raw:
            ref_cds_ph_ext = extend_cds_with_3utr(ref_cds_raw, genome, chrom, strand, settings.ALT_DOWNSTREAM_EXTENSION_BP)
            ref_cds_seq = get_coding_sequence(ref_cds_ph_ext, genome, chrom, strand)
            ref_prot_full = str(Seq(ref_cds_seq).translate(to_stop=False)) if ref_cds_seq else ""
            sp = ref_prot_full.find('*')
            ref_prot = ref_prot_full[:sp+1] if sp >= 0 else ref_prot_full
            ref_cds_ph_trunc, ref_prot_trunc = truncate_cds_at_stop(ref_cds_ph_ext, ref_prot, strand)
            ref_cds_seq_trunc = get_coding_sequence(ref_cds_ph_trunc, genome, chrom, strand)

            res['Ref_Full_CDS'] = ref_cds_seq_trunc
            res['Ref_Full_Protein'] = ref_prot_trunc
            res['Ref_CDS_Length'] = len(ref_cds_seq_trunc)
            start_codon_pos = _get_start_codon_position(ref_cds_ph_trunc, strand)

            alt_cds_ph_base = compute_alt_from_ref(ref_cds_raw, row, etype, strand)
            alt_cds_ph_ext = extend_cds_with_3utr(alt_cds_ph_base, genome, chrom, strand, settings.ALT_DOWNSTREAM_EXTENSION_BP)
            alt_cds_seq = get_coding_sequence(alt_cds_ph_ext, genome, chrom, strand)

            if not _starts_with_atg(alt_cds_seq):
                res['Alt_Full_CDS'] = ''
                res['Alt_Full_Protein'] = ''
                res['Alt_CDS_Length'] = 0
                res['Diff_CDS'] = ''
                res['Diff_Protein'] = ''
                res['Diff_Length'] = ''
                res['NMD_Status'] = 'NO_START_CODON'
                base_status = diag_msg
                res['Match_Status'] = f"{base_status};ALT_NO_START_CODON"
            else:
                alt_prot_full = str(Seq(alt_cds_seq).translate(to_stop=False)) if alt_cds_seq else ""
                sp = alt_prot_full.find('*')
                alt_prot = alt_prot_full[:sp+1] if sp >= 0 else alt_prot_full
                alt_cds_ph_trunc, alt_prot_trunc = truncate_cds_at_stop(alt_cds_ph_ext, alt_prot, strand)
                alt_cds_seq_trunc = get_coding_sequence(alt_cds_ph_trunc, genome, chrom, strand)

                res['Alt_Full_CDS'] = alt_cds_seq_trunc
                res['Alt_Full_Protein'] = alt_prot_trunc
                res['Alt_CDS_Length'] = len(alt_cds_seq_trunc)

                diff_cds = _compute_trimmed_diff(ref_cds_seq_trunc, alt_cds_seq_trunc)
                res['Diff_CDS'] = diff_cds
                res['Diff_Protein'] = _compute_trimmed_diff(ref_prot_trunc, alt_prot_trunc)
                diff_cds_len_display = len(diff_cds.replace('|', '')) if diff_cds else 0
                res['Diff_Length'] = diff_cds_len_display if diff_cds else ''

                frameshift = _is_frameshift_by_event(etype, row, ref_cds_ph=ref_cds_ph_trunc, strand=strand)

                res['NMD_Status'] = judge_nmd_status(
                    ref_prot_trunc,
                    alt_prot_trunc,
                    frameshift,
                    alt_cds_ph_trunc,
                    full_tx_exons,
                    strand,
                    start_codon_pos
                )

            if gb_output_dir:
                generate_exon_gb(ref_tid, exon_bounds, ref_cds_ph_trunc, genome, chrom, strand,
                                 display_gene_symbol, etype, row, gb_output_dir, job_id, event_id=event_id)
                generate_tran_gb(ref_tid, full_tx_exons, ref_cds_ph_trunc, genome, chrom, strand,
                                 display_gene_symbol, etype, row, gb_output_dir, job_id, event_id=event_id)

        else:
            res['Ref_Full_CDS'] = ''
            res['Ref_Full_Protein'] = ''
            res['Ref_CDS_Length'] = 0
            res['Alt_Full_CDS'] = ''
            res['Alt_Full_Protein'] = ''
            res['Alt_CDS_Length'] = 0
            res['Diff_CDS'] = ''
            res['Diff_Protein'] = ''
            res['Diff_Length'] = ''
            res['NMD_Status'] = 'NON_CODING'
            res['Match_Status'] = diag_msg

            if gb_output_dir:
                generate_exon_gb(ref_tid, exon_bounds, [], genome, chrom, strand,
                                 display_gene_symbol, etype, row, gb_output_dir, job_id, event_id=event_id)
                generate_tran_gb(ref_tid, full_tx_exons, [], genome, chrom, strand,
                                 display_gene_symbol, etype, row, gb_output_dir, job_id, event_id=event_id)

    except KeyError as e:
        res['Match_Status'] = f"[LEVEL0] ERROR: Missing column {e}"
    except IndexError as e:
        res['Match_Status'] = f"[LEVEL0] ERROR: Index out of range {e}"
    except Exception as e:
        res['Match_Status'] = f"[LEVEL0] ERROR: {type(e).__name__}: {e}"
        logger.debug(f"Event error: {e}", exc_info=True)
    return res


def _apply_left_alignment(writer, sheet_name):
    ws = writer.sheets[sheet_name]
    la = Alignment(horizontal='left', vertical='center', wrap_text=False)
    for row_cells in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row_cells:
            cell.alignment = la


def filter_and_summarize(rmats_dir, output_dir, genome_file, gtf_file,
                         fdr=0.05, psi=0.2, minjc=15, job_id=None, generate_gb=True,
                         use_cache=True, cache_dir=None):
    """Main pipeline: filter significant events -> annotate reference/alt CDS,
    protein and NMD status -> write the summary Excel (+ GenBank files)."""
    if job_id is None:
        job_id = str(uuid.uuid4())[:8]

    if cache_dir is None:
        cache_dir = settings.get_cache_dir()

    output_dir = os.path.join(output_dir, f"Job_{job_id}")
    os.makedirs(output_dir, exist_ok=True)
    output_excel = os.path.join(output_dir, f"Job_{job_id}_results.xlsx")
    start_time = datetime.now()

    logger.info(console.section(f"Job {job_id}"))
    logger.info(f"Output directory : {output_dir}")

    logger.info(console.section("Reference data"))
    logger.info(f"Genome FASTA : {genome_file}")
    genome = Fasta(genome_file, as_raw=True, sequence_always_upper=True)
    logger.info(f"GTF          : {gtf_file}")
    if use_cache:
        gene_idx, gene_exons, tid_strand, alias_map, longest_tx_per_gene = build_gtf_index_with_cache(gtf_file, cache_dir)
    else:
        gene_idx, gene_exons, tid_strand, alias_map, longest_tx_per_gene = build_gtf_index(gtf_file)
    gtf_load_time = (datetime.now() - start_time).total_seconds()
    logger.info(f"✓ {len(genome.keys()):,} genome sequences ({len(genome):,} bp) · "
                f"GTF index ready ({gtf_load_time:.1f}s)")

    gb_output_dir = None
    if generate_gb:
        gb_output_dir = os.path.join(output_dir, "GB_Files")
        os.makedirs(gb_output_dir, exist_ok=True)
        logger.info(f"GenBank output : {gb_output_dir}")

    event_files = settings.RMATS_EVENT_FILES

    all_filtered, sheet_dict = [], {}
    global_event_counter = 0

    for etype, fname in event_files.items():
        fpath = os.path.join(rmats_dir, fname)
        if not os.path.exists(fpath):
            logger.warning(f"[skip] {fname} not found in the rMATS directory")
            continue

        df = pd.read_csv(fpath, sep='\t')
        if len(df) == 0:
            logger.warning(f"[skip] {fname} contains no events")
            continue

        logger.info(console.section(f"{etype} · {fname}"))
        logger.info(f"Events loaded : {len(df):,}")

        df = _fill_missing_gene_symbols(df)

        mask = pd.Series(True, index=df.index)
        if 'FDR' in df.columns:
            mask &= df['FDR'] <= fdr
        if 'IncLevelDifference' in df.columns:
            mask &= df['IncLevelDifference'].abs() > psi

        jc_cols = ['IJC_SAMPLE_1', 'SJC_SAMPLE_1', 'IJC_SAMPLE_2', 'SJC_SAMPLE_2']
        if all(c in df.columns for c in jc_cols):
            ijc1_avg = df['IJC_SAMPLE_1'].map(parse_jc_single)
            sjc1_avg = df['SJC_SAMPLE_1'].map(parse_jc_single)
            ijc2_avg = df['IJC_SAMPLE_2'].map(parse_jc_single)
            sjc2_avg = df['SJC_SAMPLE_2'].map(parse_jc_single)

            exclude_s1 = (ijc1_avg < minjc) & (sjc1_avg < minjc)
            exclude_s2 = (ijc2_avg < minjc) & (sjc2_avg < minjc)
            exclude_ijc = (ijc1_avg < minjc) & (ijc2_avg < minjc)
            exclude_sjc = (sjc1_avg < minjc) & (sjc2_avg < minjc)

            mask &= ~(exclude_s1 | exclude_s2 | exclude_ijc | exclude_sjc)

        df_filt = df[mask].copy()
        logger.info(f"After filters : {len(df_filt):,}/{len(df):,}")
        if len(df_filt) == 0:
            logger.info("No events passed the filters")
            continue

        n_current = len(df_filt)
        current_ids = list(range(global_event_counter + 1, global_event_counter + n_current + 1))

        logger.info(f"Annotating {n_current:,} events (ID {current_ids[0]}–{current_ids[-1]}) ...")
        records = []
        event_start_time = datetime.now()

        for idx, (eid, (_, row)) in enumerate(zip(current_ids, df_filt.iterrows())):
            rec = process_event(row, etype, genome, gene_idx, gene_exons, tid_strand,
                                alias_map, longest_tx_per_gene, gb_output_dir, event_id=str(eid), job_id=job_id)
            records.append(rec)

            if (idx + 1) % 100 == 0:
                elapsed = (datetime.now() - event_start_time).total_seconds()
                rate = (idx + 1) / elapsed if elapsed > 0 else 0.0
                logger.info(f"  · {idx + 1}/{n_current:,} annotated ({elapsed:.1f}s, {rate:.0f} events/s)")

        event_time = (datetime.now() - event_start_time).total_seconds()
        logger.info(f"✓ Annotated {n_current:,} events in {event_time:.1f}s")

        annot_df = pd.DataFrame(records)[ANNOTATION_COLUMNS]

        all_sig, _ = _build_all_significant_row(df_filt.reset_index(drop=True), annot_df, etype,
                                                global_counter_start=global_event_counter)
        all_filtered.append(all_sig)

        sheet_df = df_filt.reset_index(drop=True).copy()
        sheet_df.insert(0, 'Event_ID', current_ids)
        sheet_dict[etype] = pd.concat([sheet_df, annot_df], axis=1)

        global_event_counter += n_current

        ok = sum(1 for r in records if r['Match_Status'] == 'OK')
        l2 = sum(1 for r in records if '[LEVEL2]' in r['Match_Status'])
        l1 = sum(1 for r in records if '[LEVEL1]' in r['Match_Status'])
        nsc = sum(1 for r in records if 'ALT_NO_START_CODON' in r['Match_Status'])
        noncoding = sum(1 for r in records if r.get('NMD_Status', '') == 'NON_CODING')
        nmd = defaultdict(int)
        errs = defaultdict(int)
        for r in records:
            nmd[r.get('NMD_Status', 'UNKNOWN')] += 1
            ms = r['Match_Status']
            if ms not in ('OK',) and 'ALT_NO_START_CODON' not in ms and '[LEVEL2]' not in ms and '[LEVEL1]' not in ms:
                errs[ms] += 1
        logger.info(f"Match levels  : OK {ok} · LEVEL2 {l2} · LEVEL1 {l1} · NO_START {nsc} · "
                    f"NON_CODING {noncoding} (of {len(records)})")
        nmd_str = " · ".join(f"{status or '(unmatched)'} {cnt}"
                             for status, cnt in sorted(nmd.items(), key=lambda x: -x[1]))
        logger.info(f"NMD status    : {nmd_str}")
        if errs:
            for msg, cnt in sorted(errs.items(), key=lambda x: -x[1]):
                logger.warning(f"  unmatched   : {msg} ×{cnt}")

    logger.info(console.section("Summary"))
    logger.info(f"Events annotated : {global_event_counter:,}")

    if gb_output_dir and os.path.exists(gb_output_dir):
        all_gb = glob.glob(os.path.join(gb_output_dir, "*.gb"))
        exon_gb = [f for f in all_gb if 'EXON' in os.path.basename(f)]
        tran_gb = [f for f in all_gb if 'Tran' in os.path.basename(f)]
        logger.info(f"GenBank files    : {len(all_gb):,} (EXON {len(exon_gb):,} · Tran {len(tran_gb):,})")

    if all_filtered:
        final_df = pd.concat(all_filtered, ignore_index=True)
        with pd.ExcelWriter(output_excel, engine='openpyxl') as writer:
            final_df.to_excel(writer, sheet_name='All_Significant', index=False)
            for name, dfs in sheet_dict.items():
                dfs.to_excel(writer, sheet_name=name, index=False)
            _apply_left_alignment(writer, 'All_Significant')
            for name in sheet_dict:
                _apply_left_alignment(writer, name)
        total_time = (datetime.now() - start_time).total_seconds()
        logger.info(f"✅ Analysis complete in {total_time:.1f}s")
        logger.info(f"   Output Excel : {output_excel}")
        return output_excel
    else:
        logger.warning("⚠️  No significant events found — nothing to write.")
        return None
